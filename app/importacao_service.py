"""Leitura de planilhas (CSV/XLSX) para importação em lote de associados.

Só lê e normaliza: a validação usa as mesmas regras do cadastro (em main.py).
Os nomes de coluna aceitos são os mesmos da exportação, então dá para exportar,
ajustar no Excel e importar de volta (novos registros).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_LINHAS = 5000

# campo do formulário -> (título no modelo, obrigatório, apelidos aceitos no cabeçalho)
CAMPOS = {
    'matricula': ('Matrícula', True, ()),
    'nome': ('Nome', True, ('nome completo',)),
    'cpf': ('CPF', True, ()),
    'rg': ('RG', True, ()),
    'data_nascimento': ('Data de nascimento', True, ('nascimento', 'data nascimento')),
    'data_admissao': ('Data de admissão', True, ('admissao', 'data admissao')),
    'email': ('E-mail', True, ('email',)),
    'endereco': ('Endereço', True, ('endereco completo',)),
    'telefone': ('Telefone', False, ()),
    'telefone_whatsapp': ('WhatsApp', False, ('telefone whatsapp',)),
    'situacao': ('Situação', False, ()),
    'situacao_data': ('Data da situação', False, ()),
    'situacao_motivo': ('Motivo da situação', False, ()),
    'dependentes': ('Dependentes', False, ()),
}
CAMPOS_DATA = ('data_nascimento', 'data_admissao', 'situacao_data')
COR_CABECALHO = '0B2447'


def _chave(texto) -> str:
    """'Data de Admissão ' -> 'datadeadmissao' (sem acento, caixa, espaço ou pontuação)."""
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]', '', texto.lower())


_MAPA_CABECALHO = {}
for _campo, (_titulo, _, _apelidos) in CAMPOS.items():
    for _nome in (_titulo, _campo, *_apelidos):
        _MAPA_CABECALHO[_chave(_nome)] = _campo


class PlanilhaInvalida(ValueError):
    """Problema no arquivo como um todo (formato, cabeçalho, tamanho)."""


def _texto(valor, campo) -> str:
    """Converte o valor da célula em texto, desfazendo o que o Excel costuma estragar."""
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        valor = valor.date()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if isinstance(valor, int):
        texto = str(valor)
        # CPF guardado como número perde os zeros à esquerda.
        return texto.zfill(11) if campo == 'cpf' and len(texto) < 11 else texto
    return str(valor).strip()


def _data_iso(texto: str) -> str:
    """Aceita dd/mm/aaaa, dd-mm-aaaa e aaaa-mm-dd; devolve ISO (ou o texto original se não reconhecer,
    para a validação acusar 'data inválida')."""
    for formato in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y'):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return texto


def _normalizar_linha(bruta: dict) -> dict:
    linha = {campo: _texto(valor, campo) for campo, valor in bruta.items()}
    for campo in CAMPOS_DATA:
        if linha.get(campo):
            linha[campo] = _data_iso(linha[campo])
    return linha


def _linhas_xlsx(conteudo: bytes):
    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl levanta vários tipos
        raise PlanilhaInvalida('Não foi possível abrir o arquivo Excel (.xlsx).') from exc
    ws = wb['Associados'] if 'Associados' in wb.sheetnames else wb.worksheets[0]
    for linha in ws.iter_rows(values_only=True):
        yield list(linha)


def _linhas_csv(conteudo: bytes):
    for codificacao in ('utf-8-sig', 'cp1252'):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover — cp1252 decodifica quase tudo
        raise PlanilhaInvalida('Codificação do CSV não reconhecida.')
    primeira = texto.split('\n', 1)[0]
    separador = ';' if primeira.count(';') >= primeira.count(',') else ','
    yield from csv.reader(io.StringIO(texto), delimiter=separador)


def ler_planilha(conteudo: bytes, extensao: str):
    """Lê o arquivo e devolve (linhas, colunas_ignoradas).

    `linhas` é uma lista de (número da linha na planilha, dict campo -> texto),
    sem as linhas totalmente vazias."""
    extensao = extensao.lower().lstrip('.')
    if extensao == 'xlsx':
        leitor = _linhas_xlsx(conteudo)
    elif extensao == 'csv':
        leitor = _linhas_csv(conteudo)
    else:
        raise PlanilhaInvalida('Envie um arquivo .xlsx ou .csv.')

    try:
        cabecalho = next(leitor)
    except StopIteration:
        raise PlanilhaInvalida('A planilha está vazia.') from None

    colunas, ignoradas = {}, []
    for indice, titulo in enumerate(cabecalho):
        if titulo is None or str(titulo).strip() == '':
            continue
        campo = _MAPA_CABECALHO.get(_chave(titulo))
        if campo and campo not in colunas.values():
            colunas[indice] = campo
        else:
            ignoradas.append(str(titulo).strip())

    faltando = [CAMPOS[c][0] for c, (_, obrig, _) in CAMPOS.items() if obrig and c not in colunas.values()]
    if faltando:
        raise PlanilhaInvalida(
            'Colunas obrigatórias ausentes no cabeçalho: ' + ', '.join(faltando)
            + '. Use o modelo disponível nesta página.'
        )

    linhas = []
    for numero, valores in enumerate(leitor, start=2):
        bruta = {campo: valores[i] if i < len(valores) else None for i, campo in colunas.items()}
        if all(v is None or str(v).strip() == '' for v in bruta.values()):
            continue
        if len(linhas) >= MAX_LINHAS:
            raise PlanilhaInvalida(f'A planilha tem mais de {MAX_LINHAS} linhas. Divida em arquivos menores.')
        linhas.append((numero, _normalizar_linha(bruta)))

    if not linhas:
        raise PlanilhaInvalida('Nenhuma linha de dados encontrada abaixo do cabeçalho.')
    return linhas, ignoradas


_RE_DEPENDENTE = re.compile(r'^(?P<nome>.+?)\s*\((?P<parentesco>[^()]*(?:\([^()]*\)[^()]*)*)\)\s*$')


def separar_dependentes(texto: str, parentescos, padrao: str):
    """'Maria (Filho(a)); João' -> [('Maria', 'Filho(a)'), ('João', padrao)].

    Aceita o formato da exportação ("Nome (Parentesco)") ou só nomes separados por ; ou ,."""
    resultado = []
    separador = ';' if ';' in (texto or '') else ','
    for parte in (texto or '').split(separador):
        parte = parte.strip()
        if not parte:
            continue
        m = _RE_DEPENDENTE.match(parte)
        if m and m.group('parentesco').strip() in parentescos:
            resultado.append((m.group('nome').strip(), m.group('parentesco').strip()))
        else:
            resultado.append((parte, padrao))
    return resultado


def gerar_modelo_xlsx(parentescos, situacoes) -> bytes:
    """Planilha modelo: aba 'Associados' (cabeçalho + exemplo) e aba 'Instruções'."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Associados'
    ws.append([titulo for titulo, _, _ in CAMPOS.values()])
    for indice, (campo, (_, obrigatorio, _)) in enumerate(CAMPOS.items(), start=1):
        cell = ws.cell(row=1, column=indice)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=COR_CABECALHO if obrigatorio else '5B6B82')
        cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(indice)].width = 36 if campo in ('nome', 'endereco', 'dependentes') else 18
    ws.append([
        'EX-0001', 'Nome de Exemplo (apague esta linha)', '529.982.247-25', '1234567-SSP/AL',
        date(1985, 4, 12), date(2020, 3, 1), 'exemplo@email.com', 'Rua Exemplo, 100 - Maceió/AL',
        '(82) 3333-4444', '(82) 99999-8888', 'Ativo', None, None, 'Maria Exemplo (Filho(a)); João Exemplo (Cônjuge)',
    ])
    for cell in ws[2]:
        if isinstance(cell.value, date):
            cell.number_format = 'DD/MM/YYYY'
    for col in ('A', 'C', 'I', 'J'):  # evita o Excel converter matrícula/CPF/telefone em número
        for cell in ws[col]:
            cell.number_format = '@'
    ws.freeze_panes = 'A2'

    inst = wb.create_sheet('Instruções')
    linhas = [
        ('Como preencher', ''),
        ('Colunas azul-escuras', 'Obrigatórias: Matrícula, Nome, CPF, RG, datas de nascimento e admissão, E-mail e Endereço.'),
        ('Colunas cinza', 'Opcionais.'),
        ('Datas', 'Formato dd/mm/aaaa (ou células de data do Excel).'),
        ('CPF', 'Com ou sem pontuação; precisa ser válido e não pode já existir no sistema.'),
        ('Matrícula', 'Não pode já existir no sistema nem repetir na planilha.'),
        ('Situação', 'Em branco = Ativo. Valores: ' + ', '.join(situacoes) + '.'),
        ('Dependentes', 'Separe por ponto e vírgula. Formato "Nome (Parentesco)" ou só o nome. '
                        'Parentescos: ' + ', '.join(parentescos) + '.'),
        ('Fotos', 'Não são importadas; adicione depois pela edição do associado.'),
        ('Linha de exemplo', 'Apague a linha 2 da aba Associados antes de importar.'),
    ]
    for titulo, texto in linhas:
        inst.append([titulo, texto])
    inst['A1'].font = Font(bold=True, size=13)
    for row in inst.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
    inst.column_dimensions['A'].width = 22
    inst.column_dimensions['B'].width = 100

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
