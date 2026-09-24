"""Exportação de associados em planilha (CSV para Excel pt-BR e XLSX)."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (cabeçalho, função que extrai o valor, largura da coluna no XLSX)
COLUNAS = (
    ('Matrícula', lambda a: a.matricula, 14),
    ('Nome', lambda a: a.nome, 36),
    ('CPF', lambda a: a.cpf, 16),
    ('RG', lambda a: a.rg, 18),
    ('Data de nascimento', lambda a: a.data_nascimento, 14),
    ('Data de admissão', lambda a: a.data_admissao, 14),
    ('Situação', lambda a: a.situacao_rotulo, 12),
    ('Data da situação', lambda a: a.situacao_data, 14),
    ('Motivo da situação', lambda a: a.situacao_motivo, 30),
    ('Telefone', lambda a: a.telefone, 16),
    ('WhatsApp', lambda a: a.telefone_whatsapp, 16),
    ('E-mail', lambda a: a.email, 30),
    ('Endereço', lambda a: a.endereco, 40),
    ('Dependentes', lambda a: a.dependentes, 30),
)

# Textos iniciados por estes caracteres seriam interpretados como fórmula pelo Excel
# (injeção de fórmulas em CSV/planilhas).
_INICIO_FORMULA = ('=', '+', '-', '@', '\t', '\r')

COR_CABECALHO = '0B2447'


def _texto_seguro(valor: str) -> str:
    return "'" + valor if valor.startswith(_INICIO_FORMULA) else valor


def _linhas(associados):
    for a in associados:
        yield [extrair(a) for _, extrair, _ in COLUNAS]


def gerar_csv(associados) -> bytes:
    """CSV com `;` e BOM UTF-8: abre direto no Excel em português com acentos corretos."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';', lineterminator='\r\n')
    writer.writerow([titulo for titulo, _, _ in COLUNAS])
    for linha in _linhas(associados):
        writer.writerow([
            v.strftime('%d/%m/%Y') if isinstance(v, (date, datetime))
            else _texto_seguro(str(v)) if v is not None
            else ''
            for v in linha
        ])
    return buffer.getvalue().encode('utf-8-sig')


def gerar_xlsx(associados) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Associados'

    ws.append([titulo for titulo, _, _ in COLUNAS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=COR_CABECALHO)
        cell.alignment = Alignment(vertical='center')

    for linha in _linhas(associados):
        ws.append(linha)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = 'DD/MM/YYYY'
            elif isinstance(cell.value, str):
                # Força texto: o openpyxl gravaria "=..." como fórmula.
                cell.data_type = 's'

    for indice, (_, _, largura) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
