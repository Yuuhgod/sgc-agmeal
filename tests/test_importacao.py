"""Testes da importação em lote de associados (CSV/XLSX)."""

from __future__ import annotations

import csv
import io
import uuid
from datetime import date

from openpyxl import Workbook, load_workbook

from database import ACAO_ASSOCIADO_IMPORTAR, Associado, Auditoria
from importacao_service import separar_dependentes
from tests.cpf_utils import cpf_digitos_validos
from tests.test_correcoes_seguranca import _login

CABECALHO = ['Matrícula', 'Nome', 'CPF', 'RG', 'Data de nascimento', 'Data de admissão', 'E-mail', 'Endereço',
             'Situação', 'Dependentes']


def _linha(matricula=None, cpf=None, nome=None, **extra):
    base = {
        'Matrícula': matricula or f'IMP-{uuid.uuid4().hex[:8]}',
        'Nome': nome or f'Importado {uuid.uuid4().hex[:6]}',
        'CPF': cpf or cpf_digitos_validos(),
        'RG': '123-SSP', 'Data de nascimento': '15/01/1990', 'Data de admissão': '01/03/2020',
        'E-mail': 'imp@example.invalid', 'Endereço': 'Rua X, 1', 'Situação': '', 'Dependentes': '',
    }
    base.update(extra)
    return base


def _csv(linhas, sep=';', cabecalho=CABECALHO, codificacao='utf-8-sig'):
    """CSV como o Excel grava: valores com o separador ficam entre aspas."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=sep, lineterminator='\r\n')
    w.writerow(cabecalho)
    for linha in linhas:
        w.writerow([linha.get(c, '') for c in cabecalho])
    return buf.getvalue().encode(codificacao)


def _enviar(client, conteudo, nome='associados.csv'):
    return client.post('/associados/importar', data={'arquivo': (io.BytesIO(conteudo), nome)},
                       content_type='multipart/form-data')


def test_modelo_xlsx(admin_client):
    r = admin_client.get('/associados/importar/modelo')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ['Associados', 'Instruções']
    assert [c.value for c in wb['Associados'][1]][:3] == ['Matrícula', 'Nome', 'CPF']


def test_importar_csv_previa_e_confirmacao(admin_client, flask_app):
    boas = [_linha(Dependentes='Ana (Filho(a)); Beto'), _linha(Situação='Desligado')]
    ruim = _linha(CPF='11111111111')
    r = _enviar(admin_client, _csv(boas + [ruim]))
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert '2 pronta(s) para importar' in html
    assert '1 com erro' in html
    assert 'O CPF digitado é matematicamente inválido.' in html

    # Nada gravado ainda.
    with flask_app.app_context():
        assert Associado.query.filter_by(matricula=boas[0]['Matrícula']).first() is None

    r = admin_client.post('/associados/importar/confirmar')
    assert r.status_code == 302
    with flask_app.app_context():
        a = Associado.query.filter_by(matricula=boas[0]['Matrícula']).first()
        assert a.data_nascimento == date(1990, 1, 15)
        assert [(d.nome, d.parentesco) for d in a.dependentes] == [('Ana', 'Filho(a)'), ('Beto', 'Não informado')]
        b = Associado.query.filter_by(matricula=boas[1]['Matrícula']).first()
        assert b.situacao == 'desligado' and b.situacao_data == date.today()
        assert Associado.query.filter_by(matricula=ruim['Matrícula']).first() is None
        log = Auditoria.query.filter_by(acao=ACAO_ASSOCIADO_IMPORTAR).order_by(Auditoria.id.desc()).first()
        assert log.descricao.startswith('2 associado(s) importado(s)')

    # A pré-visualização é consumida: confirmar de novo não duplica.
    r = admin_client.post('/associados/importar/confirmar', follow_redirects=True)
    assert 'expirou'.encode() in r.data


def test_duplicados_no_banco_e_na_planilha(admin_client, flask_app):
    existente = _linha()
    _enviar(admin_client, _csv([existente]))
    admin_client.post('/associados/importar/confirmar')

    repetida = _linha()
    linhas = [
        _linha(matricula=existente['Matrícula']),          # já existe no banco
        repetida,
        _linha(matricula=repetida['Matrícula']),           # repete a linha 3
        _linha(cpf=repetida['CPF']),                       # CPF repetido
    ]
    html = _enviar(admin_client, _csv(linhas)).get_data(as_text=True)
    assert f'Já existe um associado com a matrícula {existente["Matrícula"]}' in html
    assert 'Matrícula repetida na planilha (linha 3).' in html
    assert 'CPF repetido na planilha (linha 3).' in html
    assert '1 pronta(s) para importar' in html


def test_importar_xlsx_com_tipos_do_excel(admin_client, flask_app):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Associados'
    ws.append(CABECALHO)
    cpf = '0' + cpf_digitos_validos()[1:]
    while True:  # CPF válido começando com 0 (o Excel o guardaria como número sem o zero)
        from main import validar_cpf
        if validar_cpf(cpf):
            break
        cpf = '0' + cpf_digitos_validos()[1:]
    matricula = 12345 + int(uuid.uuid4().int % 10**6)
    ws.append([matricula, 'Nome Excel', int(cpf), 'RG1', date(1980, 5, 6), date(2019, 7, 8),
               'x@example.invalid', 'Rua', 'Ativo', None])
    buf = io.BytesIO()
    wb.save(buf)

    html = _enviar(admin_client, buf.getvalue(), 'planilha.xlsx').get_data(as_text=True)
    assert '1 pronta(s) para importar' in html, html[:3000]
    admin_client.post('/associados/importar/confirmar')
    with flask_app.app_context():
        a = Associado.query.filter_by(matricula=str(matricula)).first()
        assert a.cpf.replace('.', '').replace('-', '') == cpf
        assert a.data_admissao == date(2019, 7, 8)


def test_csv_latin1_com_virgula(admin_client):
    conteudo = _csv([_linha(nome='José da Conceição')], sep=',', codificacao='cp1252')
    html = _enviar(admin_client, conteudo).get_data(as_text=True)
    assert 'José da Conceição' in html
    assert '1 pronta(s) para importar' in html


def test_planilha_sem_colunas_obrigatorias(admin_client):
    r = _enviar(admin_client, _csv([{'Nome': 'x'}], cabecalho=['Nome', 'Observação']))
    assert r.status_code == 302
    html = admin_client.get('/associados/importar').get_data(as_text=True)
    assert 'Colunas obrigatórias ausentes' in html


def test_formato_invalido(admin_client):
    r = _enviar(admin_client, b'MZ...', 'virus.exe')
    assert r.status_code == 302


def test_exportar_e_reimportar_planilha_do_sistema(admin_client, flask_app):
    """A planilha exportada tem as mesmas colunas aceitas pela importação."""
    nome = f'Roundtrip {uuid.uuid4().hex[:6]}'
    _enviar(admin_client, _csv([_linha(nome=nome)]))
    admin_client.post('/associados/importar/confirmar')
    xlsx = admin_client.post('/exportar_planilha', data={'formato': 'xlsx', 'nome_export': nome}).data
    html = _enviar(admin_client, xlsx, 'exportada.xlsx').get_data(as_text=True)
    # Reimportar a mesma planilha: todas as colunas são reconhecidas e o registro já existe.
    assert 'Colunas não reconhecidas' not in html
    assert 'Já existe um associado com a matrícula' in html


def test_cancelar_remove_arquivo(admin_client, flask_app):
    import main

    _enviar(admin_client, _csv([_linha()]))
    with admin_client.session_transaction() as sess:
        token = sess['importacao']['token']
    caminho = main._arquivo_importacao(token, 'csv')
    import os
    assert os.path.isfile(caminho)
    admin_client.post('/associados/importar/cancelar')
    assert not os.path.isfile(caminho)


def test_usuario_comum_nao_importa(flask_app, usuario_comum):
    c = flask_app.test_client()
    _login(c, usuario_comum)
    assert c.get('/associados/importar').status_code == 302
    assert c.get('/associados/importar/modelo').status_code == 302


def test_separar_dependentes():
    par = ('Filho(a)', 'Cônjuge', 'Não informado')
    assert separar_dependentes('Ana (Filho(a)); Bia (Cônjuge)', par, 'Não informado') == [
        ('Ana', 'Filho(a)'), ('Bia', 'Cônjuge')]
    assert separar_dependentes('Ana, Bia', par, 'Não informado') == [
        ('Ana', 'Não informado'), ('Bia', 'Não informado')]
    assert separar_dependentes('Ana (Primo)', par, 'Não informado') == [('Ana (Primo)', 'Não informado')]
