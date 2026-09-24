"""Testes dos dependentes em tabela própria (cadastro, edição, conversão do texto antigo)."""

from __future__ import annotations

import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import text

from database import Associado, Auditoria, Dependente, db
from tests.cpf_utils import cpf_digitos_validos
from tests.test_crud_associados import _payload_cadastro
from tests.test_situacao_planilha import _payload_edicao


def _com_dependentes(dados, *linhas):
    """Acrescenta linhas dep_* (nome, parentesco, nascimento, cpf) ao payload."""
    dados = dict(dados)
    for campo, indice in (('dep_nome', 0), ('dep_parentesco', 1), ('dep_nascimento', 2), ('dep_cpf', 3)):
        dados[campo] = [linha[indice] for linha in linhas]
    return dados


def _novo(nome=None):
    return _payload_cadastro(
        matricula=f'DP-{uuid.uuid4().hex[:8]}', cpf_digits=cpf_digitos_validos(),
        nome=nome or f'Titular {uuid.uuid4().hex[:6]}',
    )


def test_cadastro_com_dependentes(admin_client, flask_app):
    cpf_dep = cpf_digitos_validos()
    dados = _com_dependentes(
        _novo(),
        ('Maria Filha', 'Filho(a)', '2010-05-01', cpf_dep),
        ('', 'Não informado', '', ''),  # linha em branco é ignorada
        ('José Cônjuge', 'Cônjuge', '', ''),
    )
    assert admin_client.post('/cadastro', data=dados).status_code == 302

    with flask_app.app_context():
        a = Associado.query.filter_by(matricula=dados['matricula']).first()
        assert [(d.nome, d.parentesco) for d in a.dependentes] == [
            ('José Cônjuge', 'Cônjuge'), ('Maria Filha', 'Filho(a)'),
        ]
        maria = a.dependentes[1]
        assert maria.data_nascimento.isoformat() == '2010-05-01'
        assert maria.cpf.replace('.', '').replace('-', '') == cpf_dep


def test_dependente_invalido_mantem_linhas_no_formulario(admin_client, flask_app):
    dados = _com_dependentes(
        _novo(),
        ('Nome Mantido Dependente', 'Filho(a)', '', '11111111111'),
        ('', 'Cônjuge', '2000-01-01', ''),
    )
    r = admin_client.post('/cadastro', data=dados)
    assert r.status_code == 200
    assert 'Dependente 1: CPF inválido.'.encode() in r.data
    assert 'Dependente 2: informe o nome.'.encode() in r.data
    assert b'value="Nome Mantido Dependente"' in r.data
    with flask_app.app_context():
        assert Associado.query.filter_by(matricula=dados['matricula']).first() is None


def test_cpf_do_dependente_igual_ao_do_titular(admin_client):
    dados = _novo()
    dados = _com_dependentes(dados, ('Dependente', 'Filho(a)', '', dados['cpf']))
    r = admin_client.post('/cadastro', data=dados)
    assert 'o CPF é o mesmo do associado titular'.encode() in r.data


def test_editar_substitui_dependentes_e_audita(admin_client, flask_app):
    dados = _com_dependentes(_novo(), ('Antigo', 'Filho(a)', '', ''))
    admin_client.post('/cadastro', data=dados)
    with flask_app.app_context():
        aid = Associado.query.filter_by(matricula=dados['matricula']).first().id

    # Sem mudanças nos dependentes: auditoria não acusa alteração neles.
    admin_client.post(f'/editar/{aid}', data=_com_dependentes(
        _payload_edicao(flask_app, aid), ('Antigo', 'Filho(a)', '', ''),
    ))
    with flask_app.app_context():
        log = Auditoria.query.filter_by(entidade_id=aid, acao='associado.editar').order_by(Auditoria.id.desc()).first()
        assert 'dependentes' not in log.detalhes

    r = admin_client.post(f'/editar/{aid}', data=_com_dependentes(
        _payload_edicao(flask_app, aid), ('Novo', 'Neto(a)', '', ''),
    ))
    assert r.status_code == 302
    with flask_app.app_context():
        a = db.session.get(Associado, aid)
        assert [d.nome for d in a.dependentes] == ['Novo']
        assert Dependente.query.filter_by(nome='Antigo', associado_id=aid).count() == 0
        log = Auditoria.query.filter_by(entidade_id=aid, acao='associado.editar').order_by(Auditoria.id.desc()).first()
        assert 'dependentes: "Antigo (Filho(a))" → "Novo (Neto(a))"' in log.detalhes


def test_excluir_associado_remove_dependentes(admin_client, flask_app):
    dados = _com_dependentes(_novo(), ('Some Junto', 'Filho(a)', '', ''))
    admin_client.post('/cadastro', data=dados)
    with flask_app.app_context():
        aid = Associado.query.filter_by(matricula=dados['matricula']).first().id
    admin_client.post(f'/excluir/{aid}')
    with flask_app.app_context():
        assert Dependente.query.filter_by(associado_id=aid).count() == 0


def test_planilha_xlsx_tem_aba_de_dependentes(admin_client):
    nome = f'Titular Planilha {uuid.uuid4().hex[:6]}'
    dados = _com_dependentes(_novo(nome), ('Dep A', 'Filho(a)', '2012-03-04', ''), ('Dep B', 'Cônjuge', '', ''))
    admin_client.post('/cadastro', data=dados)

    r = admin_client.post('/exportar_planilha', data={'formato': 'xlsx', 'nome_export': nome})
    wb = load_workbook(io.BytesIO(r.data))
    assert wb['Associados']['A1'].value == 'Matrícula'
    principal = {c.value: i for i, c in enumerate(wb['Associados'][1])}
    assert wb['Associados'][2][principal['Dependentes']].value == 'Dep A (Filho(a)); Dep B (Cônjuge)'
    linhas = [[c.value for c in row] for row in wb['Dependentes'].iter_rows(min_row=2)]
    assert [(lin[1], lin[2], lin[3]) for lin in linhas] == [(nome, 'Dep A', 'Filho(a)'), (nome, 'Dep B', 'Cônjuge')]


def test_conversao_do_texto_antigo_roda_uma_vez(admin_client, flask_app):
    import main

    dados = _novo()
    admin_client.post('/cadastro', data=dados)
    with flask_app.app_context():
        a = Associado.query.filter_by(matricula=dados['matricula']).first()
        aid = a.id
        a.dependentes_texto_legado = 'Ana Souza, Bruno Souza; Carla\n  , '
        db.session.commit()
        # Simula um banco da versão antiga: a marca de conversão ainda não existe.
        with db.engine.begin() as conn:
            conn.execute(text("DELETE FROM sgc_meta WHERE chave = 'dependentes_convertidos'"))

        main._converter_dependentes_legados()
        main._converter_dependentes_legados()  # segunda vez não duplica
        db.session.expire_all()

        deps = db.session.get(Associado, aid).dependentes
        assert [(d.nome, d.parentesco) for d in deps] == [
            ('Ana Souza', 'Não informado'), ('Bruno Souza', 'Não informado'), ('Carla', 'Não informado'),
        ]
