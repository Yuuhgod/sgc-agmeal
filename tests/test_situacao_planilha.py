"""Testes da situação cadastral do associado e da exportação em planilha (CSV/XLSX)."""

from __future__ import annotations

import csv
import io
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from database import ACAO_ASSOCIADO_EXPORTAR, Associado, Auditoria, db
from tests.cpf_utils import cpf_digitos_validos
from tests.test_crud_associados import _payload_cadastro


def _cadastrar(client, flask_app, nome='Associado Situacao'):
    matricula = f'ST-{uuid.uuid4().hex[:8].upper()}'
    r = client.post('/cadastro', data=_payload_cadastro(
        matricula=matricula, cpf_digits=cpf_digitos_validos(), nome=nome,
    ))
    assert r.status_code == 302
    with flask_app.app_context():
        return Associado.query.filter_by(matricula=matricula).first().id


def _payload_edicao(flask_app, aid, **extra):
    with flask_app.app_context():
        a = db.session.get(Associado, aid)
        dados = {
            'nome': a.nome, 'matricula': a.matricula, 'cpf': a.cpf, 'rg': a.rg,
            'telefone': a.telefone or '', 'telefone_whatsapp': a.telefone_whatsapp or '',
            'endereco': a.endereco, 'email': a.email, 'dependentes': a.dependentes or '',
            'data_nascimento': a.data_nascimento.isoformat(),
            'data_admissao': a.data_admissao.isoformat(),
            'situacao': a.situacao,
        }
    dados.update(extra)
    return dados


# --- Situação cadastral -----------------------------------------------------------------

def test_novo_cadastro_entra_como_ativo(admin_client, flask_app):
    aid = _cadastrar(admin_client, flask_app)
    with flask_app.app_context():
        a = db.session.get(Associado, aid)
        assert a.situacao == 'ativo'
        assert a.situacao_data is None


def test_desligar_associado_registra_data_motivo_e_auditoria(admin_client, flask_app):
    aid = _cadastrar(admin_client, flask_app)
    r = admin_client.post(f'/editar/{aid}', data=_payload_edicao(
        flask_app, aid, situacao='desligado', situacao_data='2024-05-10',
        situacao_motivo='Pedido do associado',
    ))
    assert r.status_code == 302

    with flask_app.app_context():
        a = db.session.get(Associado, aid)
        assert a.situacao == 'desligado'
        assert a.situacao_data == date(2024, 5, 10)
        assert a.situacao_motivo == 'Pedido do associado'
        log = (Auditoria.query.filter_by(entidade='associado', entidade_id=aid, acao='associado.editar')
               .order_by(Auditoria.id.desc()).first())
        assert 'situacao: "ativo" → "desligado"' in log.detalhes


def test_inativar_sem_data_usa_hoje(admin_client, flask_app):
    aid = _cadastrar(admin_client, flask_app)
    admin_client.post(f'/editar/{aid}', data=_payload_edicao(flask_app, aid, situacao='inativo'))
    with flask_app.app_context():
        a = db.session.get(Associado, aid)
        assert a.situacao == 'inativo'
        assert a.situacao_data == date.today()


def test_reativar_limpa_data_e_motivo(admin_client, flask_app):
    aid = _cadastrar(admin_client, flask_app)
    admin_client.post(f'/editar/{aid}', data=_payload_edicao(
        flask_app, aid, situacao='desligado', situacao_motivo='x',
    ))
    admin_client.post(f'/editar/{aid}', data=_payload_edicao(
        flask_app, aid, situacao='ativo', situacao_data='2024-01-01', situacao_motivo='ignorado',
    ))
    with flask_app.app_context():
        a = db.session.get(Associado, aid)
        assert (a.situacao, a.situacao_data, a.situacao_motivo) == ('ativo', None, None)


@pytest.mark.parametrize('extra, mensagem', [
    ({'situacao': 'suspenso'}, 'Situação inválida'),
    ({'situacao': 'desligado', 'situacao_data': (date.today() + timedelta(days=1)).isoformat()},
     'não pode ser futura'),
    ({'situacao': 'desligado', 'situacao_data': '2019-01-01'}, 'anterior à data de admissão'),
])
def test_situacao_invalida_e_recusada(admin_client, flask_app, extra, mensagem):
    aid = _cadastrar(admin_client, flask_app)
    r = admin_client.post(f'/editar/{aid}', data=_payload_edicao(flask_app, aid, **extra), follow_redirects=True)
    assert mensagem.encode() in r.data
    with flask_app.app_context():
        assert db.session.get(Associado, aid).situacao == 'ativo'


def test_listar_filtra_por_situacao(admin_client, flask_app):
    nome_ativo = f'Ativo {uuid.uuid4().hex[:6]}'
    nome_deslig = f'Desligado {uuid.uuid4().hex[:6]}'
    _cadastrar(admin_client, flask_app, nome=nome_ativo)
    aid = _cadastrar(admin_client, flask_app, nome=nome_deslig)
    admin_client.post(f'/editar/{aid}', data=_payload_edicao(flask_app, aid, situacao='desligado'))

    with flask_app.app_context():
        total_desligados = Associado.query.filter_by(situacao='desligado').count()

    r = admin_client.get(f'/listar?situacao=desligado&page=1')
    assert r.status_code == 200
    if total_desligados <= 25:
        assert nome_deslig.encode() in r.data
    assert nome_ativo.encode() not in r.data


def test_buscar_filtra_por_situacao(admin_client, flask_app):
    nome = f'Busca Situacao {uuid.uuid4().hex[:6]}'
    aid = _cadastrar(admin_client, flask_app, nome=nome)
    celula = f'<td>{nome}</td>'.encode()  # o campo de busca também mostra o nome pesquisado
    r = admin_client.post('/buscar', data={'nome': nome, 'situacao': 'desligado'})
    assert celula not in r.data

    admin_client.post(f'/editar/{aid}', data=_payload_edicao(flask_app, aid, situacao='desligado'))
    r = admin_client.post('/buscar', data={'nome': nome, 'situacao': 'desligado'})
    assert celula in r.data
    assert b'selected>Desligado' in r.data


def test_dashboard_mostra_ativos(admin_client, flask_app):
    with flask_app.app_context():
        ativos = Associado.query.filter_by(situacao='ativo').count()
    r = admin_client.get('/')
    assert r.status_code == 200
    assert b'Associados ativos' in r.data
    assert f'>{ativos}</h1>'.encode() in r.data


def test_ficha_pdf_com_situacao(admin_client, flask_app):
    aid = _cadastrar(admin_client, flask_app)
    admin_client.post(f'/editar/{aid}', data=_payload_edicao(flask_app, aid, situacao='desligado'))
    with flask_app.app_context():
        matricula = db.session.get(Associado, aid).matricula
    r = admin_client.get(f'/exportar_ficha/{matricula}')
    assert r.status_code == 200
    assert r.data.startswith(b'%PDF')


# --- Planilhas --------------------------------------------------------------------------

def test_exportar_xlsx_com_filtros(admin_client, flask_app):
    nome = f'Planilha {uuid.uuid4().hex[:6]}'
    aid = _cadastrar(admin_client, flask_app, nome=nome)
    admin_client.post(f'/editar/{aid}', data=_payload_edicao(
        flask_app, aid, situacao='inativo', situacao_data='2024-02-03',
    ))

    r = admin_client.post('/exportar_planilha', data={'formato': 'xlsx', 'nome_export': nome})
    assert r.status_code == 200
    assert 'spreadsheetml' in r.headers['Content-Type']
    assert 'attachment' in r.headers['Content-Disposition']

    ws = load_workbook(io.BytesIO(r.data)).active
    cabecalho = [c.value for c in ws[1]]
    assert cabecalho[:3] == ['Matrícula', 'Nome', 'CPF']
    assert ws.max_row == 2
    linha = dict(zip(cabecalho, [c.value for c in ws[2]]))
    assert linha['Nome'] == nome
    assert linha['Situação'] == 'Inativo'
    assert linha['Data da situação'].date() == date(2024, 2, 3)

    with flask_app.app_context():
        log = Auditoria.query.filter_by(acao=ACAO_ASSOCIADO_EXPORTAR).order_by(Auditoria.id.desc()).first()
        assert log.descricao == '1 associado(s) em XLSX'
        assert f'nome={nome}' in log.detalhes


def test_exportar_csv_abre_no_excel_e_bloqueia_formulas(admin_client, flask_app):
    nome = f'=HYPERLINK("http://x","{uuid.uuid4().hex[:6]}")'
    _cadastrar(admin_client, flask_app, nome=nome)

    r = admin_client.post('/exportar_planilha', data={'formato': 'csv', 'nome_export': nome})
    assert r.status_code == 200
    assert r.data.startswith('﻿'.encode('utf-8'))  # BOM: acentos corretos no Excel

    linhas = list(csv.reader(io.StringIO(r.data.decode('utf-8-sig')), delimiter=';'))
    assert linhas[0][:2] == ['Matrícula', 'Nome']
    assert len(linhas) == 2
    assert linhas[1][1] == "'" + nome

    r = admin_client.post('/exportar_planilha', data={'formato': 'xlsx', 'nome_export': nome})
    celula = load_workbook(io.BytesIO(r.data)).active['B2']
    assert celula.data_type == 's'
    assert celula.value == nome


def test_exportar_planilha_formato_invalido(admin_client):
    r = admin_client.post('/exportar_planilha', data={'formato': 'exe'})
    assert r.status_code == 302


def test_exportar_planilha_sem_resultados(admin_client):
    r = admin_client.post('/exportar_planilha', data={'formato': 'csv', 'nome_export': 'nao-existe-' + uuid.uuid4().hex})
    assert r.status_code == 302


def test_exportar_planilha_exige_login(client):
    r = client.post('/exportar_planilha', data={'formato': 'csv'})
    assert r.status_code == 302
    assert 'login' in (r.headers.get('Location') or '').lower()
